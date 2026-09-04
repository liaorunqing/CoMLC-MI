"""Synchronize the IEEE Access checklist with the final submission package."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "- Resubmission-Checklist-6.26.23.docx"
PATH = ROOT / "revision" / "resubmission_checklist_completed.docx"


def update_document(document: Document) -> None:
    status = (
        "SUBMISSION PACKAGE COMPLETE PENDING ETHICS CONFIRMATION — 4 SEPTEMBER 2026: "
        "analysis, manuscript with integrated appendices, figures, author photographs, "
        "response, highlighted PDF, and clean source have been assembled and checked. "
        "Before portal submission, the corresponding author must confirm the institution-specific "
        "secondary-analysis ethics wording and perform the final desktop Word review."
    )
    if len(document.paragraphs) < 2:
        raise ValueError("Unexpected checklist structure")
    document.paragraphs[1].text = status

    table = document.tables[0]
    expected = {
        1: "☑\nReviewed; 14 concerns addressed",
        2: "☑\nLanguage and terminology audited",
        3: "☑\nMath notation audited",
        4: "☑\n30 references verified; links checked",
        5: "N/A\nNo video",
        6: "☑\n14 Concern–Response–Action items",
        7: "☑\nYellow highlighted PDF prepared",
        8: "☑\nAppendices A–H integrated; no separate supplement",
        9: "☑\nBiographies + 3 photos included",
        10: "☑\n29-page manuscript PDF",
    }
    for row_index, value in expected.items():
        table.rows[row_index].cells[1].text = value
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    if not header_properties.findall(qn("w:tblHeader")):
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        header_properties.append(header)


def main() -> None:
    document = Document(SOURCE)
    update_document(document)
    document.save(PATH)
    print(PATH.resolve())


if __name__ == "__main__":
    main()
