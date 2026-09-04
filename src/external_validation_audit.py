"""Audit the public Hungarian AMI registry extract for external-validation use.

The script writes aggregate diagnostics only. It does not copy or export
patient-level registry records.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


PLACEHOLDERS = {"Unknown", "Not filled in", "Not interpretable"}

FEATURE_MAPPING = [
    {
        "harmonized_feature": "age_years",
        "hungarian_field": "Age at admission",
        "krasnoyarsk_field": "AGE",
        "hungarian_transform": "numeric age at index admission",
        "krasnoyarsk_transform": "numeric age at admission",
        "analysis_role": "strict primary",
        "timing": "admission",
        "caution": "Hungarian age contains fractional years; Krasnoyarsk age is integer years.",
    },
    {
        "harmonized_feature": "male",
        "hungarian_field": "Gender",
        "krasnoyarsk_field": "SEX",
        "hungarian_transform": "Man=1; Woman=0",
        "krasnoyarsk_transform": "male=1; female=0",
        "analysis_role": "strict primary",
        "timing": "admission",
        "caution": "Binary sex coding only; no other categories are present in either extract.",
    },
    {
        "harmonized_feature": "prior_mi",
        "hungarian_field": "History of myocardial infarction",
        "krasnoyarsk_field": "INF_ANAM",
        "hungarian_transform": "Yes=1; Not=0; Unknown/Not filled in=missing",
        "krasnoyarsk_transform": "INF_ANAM>0; missing retained",
        "analysis_role": "strict primary",
        "timing": "pre-admission history",
        "caution": "Hungarian field is binary; Krasnoyarsk records the number of prior infarctions.",
    },
    {
        "harmonized_feature": "prior_hf",
        "hungarian_field": "History of heart failure",
        "krasnoyarsk_field": "ZSN_A",
        "hungarian_transform": "Yes=1; Not=0; Unknown/Not filled in=missing",
        "krasnoyarsk_transform": "ZSN_A>0; missing retained",
        "analysis_role": "strict primary",
        "timing": "pre-admission history",
        "caution": "Krasnoyarsk stages chronic HF; Hungarian field is binary.",
    },
    {
        "harmonized_feature": "hypertension",
        "hungarian_field": "Hypertension diagnosed in history or during treatment",
        "krasnoyarsk_field": "GB + SIM_GIPERT",
        "hungarian_transform": "Yes=1; Not=0; Unknown/Not filled in=missing",
        "krasnoyarsk_transform": "GB>0 or SIM_GIPERT=1; missing retained when status is unresolved",
        "analysis_role": "expanded sensitivity",
        "timing": "not guaranteed at admission",
        "caution": "Hungarian definition permits diagnosis during treatment and is not strictly time-zero.",
    },
    {
        "harmonized_feature": "diabetes",
        "hungarian_field": "Diabetes diagnosed in history or during treatment",
        "krasnoyarsk_field": "endocr_01",
        "hungarian_transform": "Yes=1; Not=0; Unknown/Not filled in=missing",
        "krasnoyarsk_transform": "0/1 history of diabetes; missing retained",
        "analysis_role": "expanded sensitivity",
        "timing": "not guaranteed at admission",
        "caution": "Hungarian definition permits diagnosis during treatment; ascertainment differs.",
    },
    {
        "harmonized_feature": "symptom_to_admission_delay",
        "hungarian_field": "Date of occurrence of complaints + First hospital admission of an event",
        "krasnoyarsk_field": "TIME_B_S",
        "hungarian_transform": "derive elapsed hours when complaint time is known",
        "krasnoyarsk_transform": "ordinal bins 1-9",
        "analysis_role": "optional sensitivity",
        "timing": "admission",
        "caution": "Complaint onset is unavailable for about 43% of Hungarian events; bin edges require harmonization.",
    },
]

OUTCOME_MAPPING = [
    {
        "hungarian_endpoint": "death_1d",
        "definition": "death date from admission calendar date through day 1",
        "closest_krasnoyarsk_label": "LET_IS > 0",
        "compatibility": "adjacent, not equivalent",
    },
    {
        "hungarian_endpoint": "death_7d",
        "definition": "death date from admission calendar date through day 7",
        "closest_krasnoyarsk_label": "LET_IS > 0",
        "compatibility": "short-term proxy; discharge status unavailable",
    },
    {
        "hungarian_endpoint": "death_30d",
        "definition": "death date from admission calendar date through day 30",
        "closest_krasnoyarsk_label": "LET_IS > 0",
        "compatibility": "clinically relevant external fatal outcome; not in-hospital mortality",
    },
    {
        "hungarian_endpoint": "death_365d",
        "definition": "death date from admission calendar date through day 365",
        "closest_krasnoyarsk_label": "LET_IS > 0",
        "compatibility": "different prediction horizon",
    },
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def effective_missing(series: pd.Series) -> pd.Series:
    values = series.astype("string").str.strip()
    return series.isna() | values.isin(PLACEHOLDERS)


def load_hungarian(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path)
    df.columns = df.columns.astype(str).str.strip()
    return df


def patient_key_audit(df: pd.DataFrame) -> pd.DataFrame:
    grouped = df.groupby("Study ID", dropna=False)
    audit = grouped.agg(
        rows=("Event ID", "size"),
        gender_values=("Gender", "nunique"),
        age_min=("Age at admission", "min"),
        age_max=("Age at admission", "max"),
        alive_values=("Is the patient currently alive?", "nunique"),
        death_date_values=("Date of death", "nunique"),
        event_number_values=("Event number", "nunique"),
        event_number_max=("Event number", "max"),
        total_event_values=("Number of patient events", "nunique"),
    )
    summary = [
        ("all event rows", len(df)),
        ("unique Event ID", df["Event ID"].nunique()),
        ("unique Treatment ID", df["Treatment ID"].nunique()),
        ("unique Study ID", df["Study ID"].nunique()),
        ("rows with Event number = 1", int(df["Event number"].eq(1).sum())),
        ("Study ID groups with >1 row", int(audit["rows"].gt(1).sum())),
        ("Study ID groups with inconsistent gender", int(audit["gender_values"].gt(1).sum())),
        ("Study ID groups with age span >2 years", int((audit["age_max"] - audit["age_min"]).gt(2).sum())),
        ("Study ID groups with inconsistent vital status", int(audit["alive_values"].gt(1).sum())),
        ("Study ID groups with inconsistent death date", int(audit["death_date_values"].gt(1).sum())),
        ("Study ID groups with repeated/missing event numbers", int(audit["rows"].ne(audit["event_number_values"]).sum())),
    ]
    return pd.DataFrame(summary, columns=["check", "value"])


def profile_fields(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for col in df.columns:
        miss = effective_missing(df[col])
        top = df[col].astype("string").fillna("<NA>").value_counts().head(5)
        rows.append(
            {
                "field": col,
                "dtype": str(df[col].dtype),
                "rows": len(df),
                "native_missing_n": int(df[col].isna().sum()),
                "placeholder_missing_n": int((miss & df[col].notna()).sum()),
                "effective_missing_n": int(miss.sum()),
                "effective_missing_pct": float(miss.mean()),
                "unique_including_missing": int(df[col].nunique(dropna=False)),
                "top_values": " | ".join(f"{k}: {int(v)}" for k, v in top.items()),
            }
        )
    return pd.DataFrame(rows)


def first_event_outcomes(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    first = df.loc[df["Event number"].eq(1)].copy()
    admission = pd.to_datetime(first["First hospital admission of an event"], errors="coerce")
    death = pd.to_datetime(first["Date of death"], errors="coerce")
    days = (death.dt.normalize() - admission.dt.normalize()).dt.days
    endpoints = {days_limit: days.between(0, days_limit) for days_limit in (0, 1, 7, 14, 30, 90, 180, 365)}
    outcome_rows = [
        ("first-event rows", len(first), None),
        ("unique Event ID", first["Event ID"].nunique(), None),
        ("unique Study ID (not a valid patient count)", first["Study ID"].nunique(), None),
        ("recorded dead at extraction", int(first["Is the patient currently alive?"].eq("Not").sum()), float(first["Is the patient currently alive?"].eq("Not").mean())),
        ("valid death date", int(death.notna().sum()), float(death.notna().mean())),
        ("death date before admission", int(days.lt(0).sum()), float(days.lt(0).mean())),
    ]
    for limit, mask in endpoints.items():
        outcome_rows.append((f"death through day {limit}", int(mask.sum()), float(mask.mean())))
    outcomes = pd.DataFrame(outcome_rows, columns=["measure", "n", "proportion"])

    timeline = pd.DataFrame(
        [
            ("first admission minimum", admission.min()),
            ("first admission maximum", admission.max()),
            ("death date minimum", death.min()),
            ("death date maximum", death.max()),
        ],
        columns=["measure", "date"],
    )
    return outcomes, timeline


def missingness_outcome_association(df: pd.DataFrame) -> pd.DataFrame:
    first = df.loc[df["Event number"].eq(1)].copy()
    admission = pd.to_datetime(first["First hospital admission of an event"], errors="coerce")
    death = pd.to_datetime(first["Date of death"], errors="coerce")
    death_30d = (death.dt.normalize() - admission.dt.normalize()).dt.days.between(0, 30)
    fields = [
        "History of myocardial infarction",
        "History of heart failure",
        "Hypertension diagnosed in history or during treatment",
        "Diabetes diagnosed in history or during treatment",
    ]
    rows = []
    for field in fields:
        missing = effective_missing(first[field])
        for status, mask in (("observed", ~missing), ("unknown_or_not_filled", missing)):
            rows.append(
                {
                    "field": field,
                    "status": status,
                    "n": int(mask.sum()),
                    "death_30d_n": int(death_30d.loc[mask].sum()),
                    "death_30d_rate": float(death_30d.loc[mask].mean()),
                }
            )
    any_missing = pd.concat([effective_missing(first[field]).rename(field) for field in fields], axis=1).any(axis=1)
    for status, mask in (("all_four_observed", ~any_missing), ("any_field_unknown_or_not_filled", any_missing)):
        rows.append(
            {
                "field": "four mapped comorbidity fields",
                "status": status,
                "n": int(mask.sum()),
                "death_30d_n": int(death_30d.loc[mask].sum()),
                "death_30d_rate": float(death_30d.loc[mask].mean()),
            }
        )
    return pd.DataFrame(rows)


def harmonized_frames(uci: pd.DataFrame, hungarian: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    h = hungarian.loc[hungarian["Event number"].eq(1)].copy()

    u_htn = pd.Series(np.nan, index=uci.index, dtype=float)
    u_htn.loc[(uci["GB"] > 0) | uci["SIM_GIPERT"].eq(1)] = 1
    u_htn.loc[uci["GB"].eq(0) & uci["SIM_GIPERT"].eq(0)] = 0

    u = pd.DataFrame(
        {
            "age_years": uci["AGE"],
            "male": uci["SEX"],
            "prior_mi": uci["INF_ANAM"].gt(0).where(uci["INF_ANAM"].notna()),
            "prior_hf": uci["ZSN_A"].gt(0).where(uci["ZSN_A"].notna()),
            "hypertension": u_htn,
            "diabetes": uci["endocr_01"],
        }
    )
    hh = pd.DataFrame(
        {
            "age_years": h["Age at admission"],
            "male": h["Gender"].map({"Man": 1, "Woman": 0}),
            "prior_mi": h["History of myocardial infarction"].map({"Yes": 1, "Not": 0}),
            "prior_hf": h["History of heart failure"].map({"Yes": 1, "Not": 0}),
            "hypertension": h["Hypertension diagnosed in history or during treatment"].map({"Yes": 1, "Not": 0}),
            "diabetes": h["Diabetes diagnosed in history or during treatment"].map({"Yes": 1, "Not": 0}),
        }
    )
    return u.astype(float), hh.astype(float)


def pooled_smd(a: pd.Series, b: pd.Series) -> float:
    a = a.dropna().astype(float)
    b = b.dropna().astype(float)
    denom = np.sqrt((a.var(ddof=1) + b.var(ddof=1)) / 2)
    return float((b.mean() - a.mean()) / denom) if denom > 0 else np.nan


def compare_cohorts(uci_x: pd.DataFrame, hungarian_x: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for col in uci_x.columns:
        a, b = uci_x[col], hungarian_x[col]
        rows.append(
            {
                "feature": col,
                "krasnoyarsk_n_observed": int(a.notna().sum()),
                "krasnoyarsk_missing_pct": float(a.isna().mean()),
                "krasnoyarsk_mean_or_prevalence": float(a.mean()),
                "krasnoyarsk_sd": float(a.std(ddof=1)),
                "hungarian_n_observed": int(b.notna().sum()),
                "hungarian_missing_pct": float(b.isna().mean()),
                "hungarian_mean_or_prevalence": float(b.mean()),
                "hungarian_sd": float(b.std(ddof=1)),
                "standardized_mean_difference_hungary_minus_krasnoyarsk": pooled_smd(a, b),
            }
        )
    return pd.DataFrame(rows)


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_cells in sheet.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200])
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 48)
        sheet.sheet_view.showGridLines = False
    workbook.save(path)


def write_outputs(input_path: Path, uci_path: Path, output_dir: Path) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    h = load_hungarian(input_path)
    uci = pd.read_csv(uci_path)

    field_profile = profile_fields(h)
    key_audit = patient_key_audit(h)
    outcomes, timeline = first_event_outcomes(h)
    feature_mapping = pd.DataFrame(FEATURE_MAPPING)
    outcome_mapping = pd.DataFrame(OUTCOME_MAPPING)
    missingness_outcome = missingness_outcome_association(h)
    uci_x, h_x = harmonized_frames(uci, h)
    cohort_comparison = compare_cohorts(uci_x, h_x)

    issues = pd.DataFrame(
        [
            ("P0", "Study ID is not a patient key", "Repeated Study ID values can have incompatible age, sex, vital status and death date.", "Use Event number=1 for an independent first-event cohort; do not deduplicate or cluster by Study ID."),
            ("P0", "Outcome definition mismatch", "Krasnoyarsk LET_IS is an in-hospital lethal outcome; the Hungarian extract supplies vital status and date of death but no discharge date.", "Call the analysis partial fatal-outcome transportability; report 7-day and 30-day mortality separately."),
            ("P1", "Only one overlapping outcome", "The Hungarian extract has no AF, SVT, VT, VF, AV block, edema, rupture, Dressler syndrome, CHF, reinfarction or post-infarction angina outcomes.", "Do not claim external validation of the 12-label framework."),
            ("P1", "Limited strict time-zero feature overlap", "Only age, sex, prior MI and prior HF map cleanly at admission.", "Use a prespecified strict four-feature model; treat hypertension and diabetes as sensitivity features."),
            ("P1", "Missingness is outcome-associated", "Unknown/not-filled history fields have substantially higher short-term mortality.", "Report imputed and complete-case sensitivity analyses; do not silently code unknown as absence."),
            ("P1", "Metadata/publication inconsistency", "The file contains 2018-2019 admissions; the associated publication describes events through 2021 and reports 29,596 unique patients.", "Seek author clarification and disclose the file-level audit."),
        ],
        columns=["severity", "issue", "evidence", "recommended_action"],
    )

    tables = {
        "Field_Profile": field_profile,
        "Patient_Key_Audit": key_audit,
        "Outcome_Summary": outcomes,
        "Timeline": timeline,
        "Feature_Mapping": feature_mapping,
        "Outcome_Mapping": outcome_mapping,
        "Missingness_Outcome": missingness_outcome,
        "Cohort_Comparison": cohort_comparison,
        "Issues": issues,
    }
    for name, table in tables.items():
        table.to_csv(output_dir / f"{name.lower()}.csv", index=False)

    workbook_path = output_dir / "hungarian_registry_audit.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name[:31], index=False)
    style_workbook(workbook_path)

    summary = {
        "input": str(input_path.resolve()),
        "sha256": sha256(input_path),
        "sheet_rows": int(len(h)),
        "sheet_columns": int(h.shape[1]),
        "first_event_rows": int(h["Event number"].eq(1).sum()),
        "unique_event_id": int(h["Event ID"].nunique()),
        "unique_study_id": int(h["Study ID"].nunique()),
        "years": {str(k): int(v) for k, v in h["Year of event start"].value_counts().sort_index().items()},
        "strict_common_features": ["age_years", "male", "prior_mi", "prior_hf"],
        "expanded_sensitivity_features": ["hypertension", "diabetes"],
        "key_conclusion": "Suitable only for partial fatal-outcome transportability, not full 12-label external validation.",
    }
    with (output_dir / "hungarian_registry_audit.json").open("w", encoding="utf-8") as stream:
        json.dump(summary, stream, indent=2, ensure_ascii=False)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--hungarian-xlsx", required=True, type=Path)
    parser.add_argument("--uci-csv", default=Path("dataset/Myocardial infarction complications Database.csv"), type=Path)
    parser.add_argument("--output-dir", default=Path("output/external_validation"), type=Path)
    args = parser.parse_args()
    result = write_outputs(args.hungarian_xlsx, args.uci_csv, args.output_dir)
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
