"""Final consistency and packaging checks for the CoMLC-MI submission."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

import fitz
import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def require(value: bool, message: str) -> None:
    if not value:
        raise AssertionError(message)


def check(manuscript: Path, package: Path) -> dict[str, object]:
    text = manuscript.read_text(encoding="utf-8")
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", text, re.S)
    require(abstract is not None, "Abstract missing")
    words = len(re.findall(r"\b[A-Za-z0-9][A-Za-z0-9'-]*\b", abstract.group(1)))
    require(200 <= words <= 250, f"Abstract word count is {words}")
    banned = (
        "Supplementary Material",
        "Supplementary Section",
        "access-2026-35668-r1",
        "TabPFN v2",
        "ML-KNN",
        "prespecified external validation",
    )
    require(all(item not in text for item in banned), "Banned legacy manuscript wording remains")
    require("\\input{paper/appendices.tex}" in text, "Integrated appendices missing")
    require(text.count("\\bibitem{") == 30, "Expected 30 references")

    portal = package / "01_Portal_Upload"
    required = {
        "Author_Response.docx",
        "Highlighted_Manuscript.pdf",
        "Main_Manuscript.pdf",
        "Clean_LaTeX_Source.zip",
    }
    require(required.issubset({p.name for p in portal.iterdir()}), "Portal files incomplete")
    clean = fitz.open(portal / "Main_Manuscript.pdf")
    highlighted = fitz.open(portal / "Highlighted_Manuscript.pdf")
    require(len(clean) == len(highlighted) == 29, "PDF page counts differ")
    require(sum(1 for page in highlighted for _ in (page.annots() or [])) > 0, "Highlighted PDF has no annotations")
    clean.close()
    highlighted.close()

    response = Document(portal / "Author_Response.docx")
    response_text = "\n".join(p.text for p in response.paragraphs)
    require(sum(p.text.startswith("Comment ") for p in response.paragraphs) == 14, "Response does not contain 14 comments")
    require("Technical Audit Summary" in response_text, "Technical audit summary missing")

    code_zip = package / "03_Code_and_Results" / "CoMLC_MI_Code_and_Results.zip"
    with zipfile.ZipFile(code_zip) as archive:
        names = [name.lower() for name in archive.namelist()]
        require(archive.testzip() is None, "Code ZIP failed CRC")
    require(not any("external_transportability_predictions" in name for name in names), "External patient predictions included")
    require(not any("external_transportability_bootstrap" in name for name in names), "External bootstrap array included")
    require(not any(name.endswith((".xlsx", ".xls", ".xlsm", ".ckpt", ".pt", ".pth")) for name in names), "Restricted artifact included")

    estimates = pd.read_csv(ROOT / "output" / "benchmark" / "primary_paired_difference.csv").set_index("metric")
    require(round(float(estimates.loc["macro_auroc", "estimate"]), 4) == 0.0246, "Macro-AUROC point estimate changed")
    return {
        "status": "submission_complete_pending_ethics_confirmation",
        "abstract_words": words,
        "references": 30,
        "pdf_pages": 29,
        "reviewer_comments": 14,
        "privacy_boundary": "passed",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--manuscript", type=Path, default=ROOT / "paper" / "gai_revised_clean.tex")
    parser.add_argument("--package", type=Path, default=ROOT / "output" / "final_delivery")
    parser.add_argument("--output", type=Path, default=ROOT / "output" / "submission_qa.json")
    args = parser.parse_args()
    report = check(args.manuscript.resolve(), args.package.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
