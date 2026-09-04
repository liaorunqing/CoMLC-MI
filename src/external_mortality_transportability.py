"""Feasibility analysis for fatal-outcome transportability to the Hungarian AMI extract.

This is deliberately not labelled full external validation: the source outcome
is in-hospital lethal outcome, whereas the external outcomes are date-derived
short-term all-cause mortality endpoints.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd
import statsmodels.api as sm
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import average_precision_score, brier_score_loss, roc_auc_score
from sklearn.model_selection import StratifiedKFold, cross_val_predict
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler


SEED = 42
FEATURE_SETS = {
    "strict4": ["age_years", "male", "prior_mi", "prior_hf"],
    "expanded6": ["age_years", "male", "prior_mi", "prior_hf", "hypertension", "diabetes"],
}


def harmonize(uci: pd.DataFrame, registry: pd.DataFrame) -> tuple[pd.DataFrame, pd.Series, pd.DataFrame, dict[str, pd.Series]]:
    h = registry.loc[registry["Event number"].eq(1)].copy()
    u_htn = pd.Series(np.nan, index=uci.index, dtype=float)
    u_htn.loc[(uci["GB"] > 0) | uci["SIM_GIPERT"].eq(1)] = 1
    u_htn.loc[uci["GB"].eq(0) & uci["SIM_GIPERT"].eq(0)] = 0

    x_source = pd.DataFrame(
        {
            "age_years": uci["AGE"],
            "male": uci["SEX"],
            "prior_mi": uci["INF_ANAM"].gt(0).where(uci["INF_ANAM"].notna()),
            "prior_hf": uci["ZSN_A"].gt(0).where(uci["ZSN_A"].notna()),
            "hypertension": u_htn,
            "diabetes": uci["endocr_01"],
        }
    ).astype(float)
    y_source = uci["LET_IS"].fillna(0).gt(0).astype(int)

    x_external = pd.DataFrame(
        {
            "age_years": h["Age at admission"],
            "male": h["Gender"].map({"Man": 1, "Woman": 0}),
            "prior_mi": h["History of myocardial infarction"].map({"Yes": 1, "Not": 0}),
            "prior_hf": h["History of heart failure"].map({"Yes": 1, "Not": 0}),
            "hypertension": h["Hypertension diagnosed in history or during treatment"].map({"Yes": 1, "Not": 0}),
            "diabetes": h["Diabetes diagnosed in history or during treatment"].map({"Yes": 1, "Not": 0}),
        }
    ).astype(float)

    admission = pd.to_datetime(h["First hospital admission of an event"], errors="coerce")
    death = pd.to_datetime(h["Date of death"], errors="coerce")
    days = (death.dt.normalize() - admission.dt.normalize()).dt.days
    endpoints = {f"death_{limit}d": days.between(0, limit).astype(int) for limit in (1, 7, 30, 365)}
    return x_source, y_source, x_external, endpoints


def model(add_indicator: bool):
    return make_pipeline(
        SimpleImputer(strategy="median", add_indicator=add_indicator),
        StandardScaler(),
        LogisticRegression(max_iter=5000, random_state=SEED),
    )


def calibration(y: np.ndarray, p: np.ndarray) -> tuple[float, float]:
    eps = np.finfo(float).eps
    linear_predictor = np.log(np.clip(p, eps, 1 - eps) / np.clip(1 - p, eps, 1 - eps))
    design = sm.add_constant(linear_predictor)
    fitted = sm.GLM(y, design, family=sm.families.Binomial()).fit()
    return float(fitted.params[0]), float(fitted.params[1])


def bootstrap_intervals(y: np.ndarray, p: np.ndarray, n_boot: int) -> dict[str, list[float]]:
    rng = np.random.default_rng(SEED)
    positive = np.flatnonzero(y == 1)
    negative = np.flatnonzero(y == 0)
    estimates = {"auc": [], "auprc": [], "brier": []}
    for _ in range(n_boot):
        idx = np.concatenate(
            [rng.choice(positive, len(positive), replace=True), rng.choice(negative, len(negative), replace=True)]
        )
        rng.shuffle(idx)
        estimates["auc"].append(roc_auc_score(y[idx], p[idx]))
        estimates["auprc"].append(average_precision_score(y[idx], p[idx]))
        estimates["brier"].append(brier_score_loss(y[idx], p[idx]))
    return {key: [float(x) for x in np.quantile(values, [0.025, 0.975])] for key, values in estimates.items()}


def metrics(y: pd.Series, p: np.ndarray, n_boot: int) -> dict:
    yy = y.to_numpy(dtype=int)
    intervals = bootstrap_intervals(yy, p, n_boot)
    intercept, slope = calibration(yy, p)
    observed = float(yy.mean())
    expected = float(p.mean())
    return {
        "n": int(len(yy)),
        "events": int(yy.sum()),
        "prevalence": observed,
        "mean_predicted_probability": expected,
        "auc": float(roc_auc_score(yy, p)),
        "auc_ci_low": intervals["auc"][0],
        "auc_ci_high": intervals["auc"][1],
        "auprc": float(average_precision_score(yy, p)),
        "auprc_ci_low": intervals["auprc"][0],
        "auprc_ci_high": intervals["auprc"][1],
        "brier": float(brier_score_loss(yy, p)),
        "brier_ci_low": intervals["brier"][0],
        "brier_ci_high": intervals["brier"][1],
        "calibration_intercept": intercept,
        "calibration_slope": slope,
        "expected_over_observed": expected / observed,
        "observed_over_expected": observed / expected,
    }


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_cells in sheet.columns:
            max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells[:200])
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 42)
    workbook.save(path)


def run(input_path: Path, uci_path: Path, output_dir: Path, n_boot: int) -> pd.DataFrame:
    output_dir.mkdir(parents=True, exist_ok=True)
    registry = pd.read_excel(input_path)
    registry.columns = registry.columns.astype(str).str.strip()
    uci = pd.read_csv(uci_path)
    x_source, y_source, x_external, endpoints = harmonize(uci, registry)

    rows = []
    cv = StratifiedKFold(n_splits=5, shuffle=True, random_state=SEED)
    for feature_set, columns in FEATURE_SETS.items():
        for indicator_name, add_indicator in (("median_imputation", False), ("median_plus_missing_indicators", True)):
            estimator = model(add_indicator)
            internal_p = cross_val_predict(
                estimator, x_source[columns], y_source, cv=cv, method="predict_proba", n_jobs=None
            )[:, 1]
            internal = metrics(y_source, internal_p, n_boot)
            internal.update(
                {
                    "cohort": "Krasnoyarsk 5-fold out-of-fold",
                    "endpoint": "in-hospital lethal outcome",
                    "feature_set": feature_set,
                    "missing_data_strategy": indicator_name,
                    "analysis_subset": "all",
                }
            )
            rows.append(internal)

            estimator.fit(x_source[columns], y_source)
            external_p = estimator.predict_proba(x_external[columns])[:, 1]
            for endpoint, y_external in endpoints.items():
                all_result = metrics(y_external, external_p, n_boot)
                all_result.update(
                    {
                        "cohort": "Hungarian first-event cohort",
                        "endpoint": endpoint,
                        "feature_set": feature_set,
                        "missing_data_strategy": indicator_name,
                        "analysis_subset": "all with source-fitted imputation",
                    }
                )
                rows.append(all_result)

                complete = x_external[columns].notna().all(axis=1)
                complete_result = metrics(y_external.loc[complete], external_p[complete.to_numpy()], n_boot)
                complete_result.update(
                    {
                        "cohort": "Hungarian first-event cohort",
                        "endpoint": endpoint,
                        "feature_set": feature_set,
                        "missing_data_strategy": indicator_name,
                        "analysis_subset": "complete case",
                    }
                )
                rows.append(complete_result)

    result = pd.DataFrame(rows)
    ordered = [
        "cohort", "endpoint", "feature_set", "missing_data_strategy", "analysis_subset",
        "n", "events", "prevalence", "mean_predicted_probability", "auc", "auc_ci_low",
        "auc_ci_high", "auprc", "auprc_ci_low", "auprc_ci_high", "brier", "brier_ci_low",
        "brier_ci_high", "calibration_intercept", "calibration_slope",
        "expected_over_observed", "observed_over_expected",
    ]
    result = result[ordered]
    result.to_csv(output_dir / "mortality_transportability_results.csv", index=False)
    with (output_dir / "mortality_transportability_results.json").open("w", encoding="utf-8") as stream:
        json.dump(result.to_dict(orient="records"), stream, indent=2, ensure_ascii=False)
    workbook_path = output_dir / "mortality_transportability_results.xlsx"
    notes = pd.DataFrame(
        [
            ("scope", "Feasibility analysis of fatal-outcome transportability; not full 12-label external validation."),
            ("independent unit", "One first AMI event per putative patient, defined by Event number = 1; Study ID was not used as a patient key."),
            ("primary specification", "strict4 with source-fitted median imputation; complete-case analysis is a sensitivity check."),
            ("outcome mismatch", "Krasnoyarsk predicts in-hospital lethal outcome; Hungarian endpoints are calendar-date-derived all-cause mortality."),
            ("bootstrap", f"Exploratory class-stratified patient bootstrap with {n_boot} resamples; increase before manuscript submission."),
            ("missingness", "Missing-indicator models are sensitivity analyses because registry history-field missingness is strongly outcome-associated."),
        ],
        columns=["item", "note"],
    )
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="Metrics", index=False)
        notes.to_excel(writer, sheet_name="Interpretation", index=False)
    style_workbook(workbook_path)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--hungarian-xlsx", required=True, type=Path)
    parser.add_argument("--uci-csv", default=Path("dataset/Myocardial infarction complications Database.csv"), type=Path)
    parser.add_argument("--output-dir", default=Path("output/external_validation"), type=Path)
    parser.add_argument("--bootstrap", default=500, type=int)
    args = parser.parse_args()
    result = run(args.hungarian_xlsx, args.uci_csv, args.output_dir, args.bootstrap)
    print(result.to_string(index=False))


if __name__ == "__main__":
    main()
