"""Submission-level consistency gates for IEEE Access resubmission artifacts."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = ROOT / "paper" / "gai_revised_clean.tex"
RESPONSE_MD = ROOT / "revision" / "response_to_reviewers.md"
RESPONSE_DOCX = ROOT / "revision" / "response_to_reviewers.docx"
RESULTS = ROOT / "output" / "external_validation" / "external_transportability_results.csv"
BOOTSTRAP = ROOT / "output" / "external_validation" / "external_transportability_bootstrap.npz"
SUPPLEMENT_ZIP = ROOT / "output" / "submission" / "Access-2026-35668_Supplementary_Materials.zip"
SOURCE_ZIP = ROOT / "output" / "submission" / "final_upload_20260831_v3" / "Access-2026-35668_Clean_LaTeX_Source.zip"
HIGHLIGHT_AUDIT = ROOT / "output" / "pdf" / "gai_revised_highlighted_full_figure_fixed_audit.json"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def manuscript_gates(text: str) -> dict[str, object]:
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", text, re.S)
    require(abstract is not None, "Abstract not found")
    abstract_words = len(re.findall(r"\b[\w'-]+\b", abstract.group(1)))
    require(200 <= abstract_words <= 250, f"Abstract has {abstract_words} words")

    body, bibliography = text.split("\\begin{thebibliography}{00}", 1)
    first_citations: list[str] = []
    for group in re.findall(r"\\cite\{([^}]+)\}", body):
        for key in map(str.strip, group.split(",")):
            if key and key not in first_citations:
                first_citations.append(key)
    bibitems = re.findall(r"\\bibitem\{([^}]+)\}", bibliography)
    require(first_citations == bibitems, "Bibliography is not in first-citation order")
    require(len(bibitems) == 29, f"Expected 29 cited references, found {len(bibitems)}")

    stale = [
        "lacks external validation",
        "no external validation",
        "neither discrimination nor calibration is externally validated",
    ]
    require(not any(phrase in text.lower() for phrase in stale), "Unqualified legacy external-validation wording remains")
    for required in (
        "29,596",
        "3,888",
        "0.696",
        "0.705",
        "0.1093",
        "0.1106",
        "0.794",
        "0.698",
        "partial external transportability",
    ):
        require(required.lower() in text.lower(), f"Missing frozen manuscript value/phrase: {required}")
    require("\\mathrm{LC}" in text and "\\mathrm{LD}" in text and "\\mathrm{LDS}" in text and "\\mathrm{BS}" in text, "Upright multi-letter math symbols are incomplete")
    return {"abstract_words": abstract_words, "references": len(bibitems)}


def external_result_gates() -> dict[str, object]:
    results = pd.read_csv(RESULTS)
    primary = results.query("cohort == 'Hungarian AMI registry' and endpoint == 'death_30d' and subset == 'all first events'")
    require(len(primary) == 2, "Primary external result rows are incomplete")
    require(primary["n"].eq(29596).all(), "Primary external n is not 29,596")
    require(primary["events"].eq(3888).all(), "Primary external deaths are not 3,888")
    seven_day = results.query("cohort == 'Hungarian AMI registry' and endpoint == 'death_7d' and subset == 'all first events'")
    require(seven_day["events"].eq(2408).all(), "Seven-day deaths are not 2,408")

    numeric = results.select_dtypes(include=["number"]).to_numpy(dtype=float)
    require(np.isfinite(numeric).all(), "Non-finite value in external result table")
    bootstrap = np.load(BOOTSTRAP)
    require(bootstrap.files, "Bootstrap archive is empty")
    require(all(bootstrap[key].shape == (2000,) for key in bootstrap.files), "Bootstrap array length is not 2,000")
    require(all(np.isfinite(bootstrap[key]).all() for key in bootstrap.files), "Non-finite bootstrap value")
    return {"primary_n": 29596, "death_30d": 3888, "death_7d": 2408, "bootstrap_replicates": 2000}


def response_gates() -> dict[str, object]:
    markdown = RESPONSE_MD.read_text(encoding="utf-8")
    require(len(re.findall(r"^### Comment", markdown, re.M)) == 14, "Response does not contain 14 comments")
    require("**Original title:**" in markdown and "**Revised title:**" in markdown, "Response title metadata incomplete")
    require("239 words" in markdown, "Response Abstract word count is stale")
    document = Document(RESPONSE_DOCX)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    require(sum(paragraph.text.startswith("Comment ") for paragraph in document.paragraphs) == 14, "DOCX does not contain 14 comments")
    reviewers = [paragraph for paragraph in document.paragraphs if paragraph.text.startswith("Reviewer ")]
    require(len(reviewers) == 2 and all(paragraph.paragraph_format.page_break_before for paragraph in reviewers), "Reviewer page breaks are missing")
    prohibited = ["placeholder", "todo", "tbd", "[insert", "xxxx"]
    require(not any(item in text.lower() for item in prohibited), "Response DOCX contains a placeholder")
    return {"paragraphs": len(document.paragraphs), "comments": 14, "reviewer_page_breaks": 2}


def package_gates() -> dict[str, object]:
    with zipfile.ZipFile(SUPPLEMENT_ZIP) as archive:
        supplement_names = {Path(name).name for name in archive.namelist()}
    required = {
        "supplementary_external_validation.pdf",
        "external_transportability_results.csv",
        "external_transportability_paired_differences.csv",
        "external_transportability_calibration_bins.csv",
        "external_transportability_cohort_summary.csv",
        "external_transportability_metadata.json",
        "external_transportability_results.xlsx",
        "README.txt",
    }
    require(required.issubset(supplement_names), "Supplement ZIP is missing required aggregate files")
    require(not any(name.endswith(".npz") or "prediction" in name.lower() for name in supplement_names), "Supplement ZIP contains patient-level predictions/bootstrap arrays")

    with zipfile.ZipFile(SOURCE_ZIP) as archive:
        source_names = set(archive.namelist())
    require("gai_revised_clean.tex" in source_names, "Clean source ZIP lacks manuscript TeX")
    require(not any(name.startswith("build/") or name.endswith((".aux", ".log")) for name in source_names), "Clean source ZIP contains build artifacts")
    photo_names = {name for name in source_names if name.startswith("author_photos/") and name.lower().endswith((".jpg", ".jpeg", ".png"))}
    require(len(photo_names) == 3, f"Clean source ZIP should contain 3 author photographs, found {len(photo_names)}")
    require("figures/revised/fig_cooccurrence.pdf" in source_names, "Clean source ZIP lacks revised vector co-occurrence figure")
    return {
        "supplement_files": len(supplement_names),
        "source_files": len(source_names),
        "source_author_photos": len(photo_names),
        "revised_cooccurrence_figure": True,
    }


def photo_gate() -> dict[str, object]:
    manuscript = MANUSCRIPT.read_text(encoding="utf-8")
    biography_photos = re.findall(r"\\begin\{IEEEbiography\}\[([^]]+)\]", manuscript)
    require(len(biography_photos) == 3, f"Expected 3 biography photographs, found {len(biography_photos)}")
    require(all("keepaspectratio" in item for item in biography_photos), "Biography photograph aspect-ratio protection is incomplete")
    audit = json.loads(HIGHLIGHT_AUDIT.read_text(encoding="utf-8"))
    photo_outlines = audit.get("biography_photo_outlines", [])
    require(len(photo_outlines) == 3, f"Highlighted PDF should mark 3 biography photographs, found {len(photo_outlines)}")
    return {
        "status": "pass",
        "photos_embedded": len(biography_photos),
        "required": 3,
        "highlighted_photo_outlines": len(photo_outlines),
    }


def main() -> None:
    report = {
        "manuscript": manuscript_gates(MANUSCRIPT.read_text(encoding="utf-8")),
        "external_results": external_result_gates(),
        "response": response_gates(),
        "packages": package_gates(),
        "author_photos": photo_gate(),
    }
    output = ROOT / "output" / "submission" / "revision_qa_report.json"
    output.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
