"""Acceptance gates for the Access-2026-35668 R1 reconstruction.

Run the analysis gates as soon as the locked experiments finish::

    python -m src.revision_r1_qa --level analysis

Use ``--level submission`` only after the clean manuscript, highlighted PDF,
response letter, supplement, and release archive have been generated.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path

import fitz
import numpy as np
import pandas as pd
from docx import Document

from .revision_features import feature_sets
from .revision_models import MODEL_NAMES


ALL_PREDICTION_NAMES = MODEL_NAMES + [
    "Ensemble-LP-RF-TabPFN-Equal",
    "Ensemble-LP-RF-TabPFN-Weighted",
]


ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "configs" / "access_2026_35668_r1.json"
RESULTS = ROOT / "output" / "revision_r1"
MANUSCRIPT = ROOT / "paper" / "gai_revised_clean.tex"


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def _finite_csv(path: Path) -> pd.DataFrame:
    require(path.exists(), f"Missing required result: {path}")
    frame = pd.read_csv(path)
    numeric = frame.select_dtypes(include=["number"])
    require(np.isfinite(numeric.to_numpy(dtype=float)).all(), f"Non-finite value in {path.name}")
    return frame


def config_and_feature_gates() -> dict[str, object]:
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    require(config["feature_contract"] == "admission_safe_v1", "Primary feature contract changed")
    require(config["validation"]["outer_repeats"] == 5, "Outer repeats must be five")
    require(config["validation"]["outer_folds"] == 5, "Outer folds must be five")
    require(config["validation"]["inner_folds"] == 4, "Inner folds must be four")
    require(config["statistics"]["patient_bootstrap"] == 2000, "Bootstrap must use 2,000 draws")
    require(config["statistics"]["paired_permutations"] == 10000, "Rare-label permutation must use 10,000 draws")
    require(config["gcn_ablation"]["seeds"] == list(range(42, 52)), "GCN seeds must be 42--51")
    raw = pd.read_csv(ROOT / config["dataset"])
    contracts = feature_sets(raw.columns)
    expected = {
        "admission_safe_v1": 91,
        "prospective_24h": 94,
        "prospective_48h": 97,
        "prospective_72h": 100,
        "retrospective_full_111": 111,
    }
    observed = {name: len(contracts[name]) for name in expected}
    require(observed == expected, f"Feature counts changed: {observed}")
    return {"feature_counts": observed, "bootstrap": 2000, "gcn_seeds": "42--51"}


def checkpoint_gates() -> dict[str, object]:
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    internal_specification = {
        "seed": config["seed"],
        "dataset": config["dataset"],
        "feature_contract": config["feature_contract"],
        "validation": config["validation"],
        "preprocessing": config["preprocessing"],
        "models": config["models"],
        "ensembles": config["ensembles"],
    }
    expected_fingerprint = hashlib.sha256(
        json.dumps(internal_specification, sort_keys=True).encode("utf-8")
    ).hexdigest()
    checkpoint_dir = RESULTS / "checkpoints"
    metadata_paths = sorted(checkpoint_dir.glob("repeat_*_fold_*.json"))
    array_paths = sorted(checkpoint_dir.glob("repeat_*_fold_*.npz"))
    require(len(metadata_paths) == 25 and len(array_paths) == 25, "Expected 25 complete outer-fold checkpoints")
    seen_pairs: set[tuple[int, int]] = set()
    for path in metadata_paths:
        meta = json.loads(path.read_text(encoding="utf-8"))
        pair = (int(meta["repeat"]), int(meta["outer_fold"]))
        require(pair not in seen_pairs, f"Duplicate outer fold {pair}")
        seen_pairs.add(pair)
        train = set(map(int, meta["outer_fit_rows"]))
        test = set(map(int, meta["outer_test_rows"]))
        require(train.isdisjoint(test), f"Outer leakage in repeat/fold {pair}")
        require(int(meta["feature_count"]) == 91, f"Wrong feature count in repeat/fold {pair}")
        require(len(meta["tabpfn_selected_feature_indices"]) == 80, f"Wrong TabPFN subset in {pair}")
        require(meta.get("internal_analysis_sha256") == expected_fingerprint, f"Checkpoint fingerprint mismatch in {pair}")
        for inner in meta["inner_folds"]:
            require(
                set(map(int, inner["fit_rows"])).isdisjoint(map(int, inner["validation_rows"])),
                f"Inner leakage in repeat/fold {pair}",
            )
    return {"outer_checkpoints": 25, "unique_repeat_fold_pairs": len(seen_pairs)}


def internal_result_gates() -> dict[str, object]:
    archive_path = RESULTS / "internal_oof_predictions.npz"
    require(archive_path.exists(), "Missing repeated OOF prediction archive")
    archive = np.load(archive_path)
    y_true = archive["y_true"]
    require(y_true.shape == (1700, 12), f"Unexpected outcome shape {y_true.shape}")
    for name in ALL_PREDICTION_NAMES:
        key = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
        require(key in archive.files, f"Missing OOF predictions for {name}")
        value = archive[key]
        require(value.shape == (5, 1700, 12), f"Wrong repeated OOF shape for {name}: {value.shape}")
        require(np.isfinite(value).all(), f"Non-finite OOF prediction for {name}")
    summary = _finite_csv(RESULTS / "internal_bootstrap_summary.csv")
    require(summary["n_bootstrap"].eq(2000).all(), "Internal bootstrap count is not 2,000")
    require("interval_method" in summary.columns, "Bootstrap interval method is not recorded")
    ece_rows = summary[summary["metric"].eq("ece_10")]
    other_rows = summary[~summary["metric"].eq("ece_10")]
    require(
        ece_rows["interval_method"].eq("bootstrap standard-error normal").all(),
        "Aggregate ECE must use the disclosed bootstrap standard-error interval",
    )
    require(
        ((ece_rows["ci_low"] <= ece_rows["estimate"]) & (ece_rows["estimate"] <= ece_rows["ci_high"])).all(),
        "An aggregate ECE interval does not contain its point estimate",
    )
    require(
        other_rows["interval_method"].eq("bootstrap percentile").all(),
        "Non-ECE internal intervals must use bootstrap percentiles",
    )
    repeat_metrics = _finite_csv(RESULTS / "internal_repeat_metrics.csv")
    require(
        len(repeat_metrics) == len(ALL_PREDICTION_NAMES) * 5
        and repeat_metrics.groupby("model")["repeat"].nunique().eq(5).all(),
        "Per-repeat internal metrics are incomplete",
    )
    contrast = _finite_csv(RESULTS / "primary_paired_difference.csv")
    require(set(contrast["metric"]) >= {"macro_auroc", "macro_auprc", "brier", "ece_10", "macro_f1_0_5"}, "Primary contrast metrics incomplete")
    tests = _finite_csv(RESULTS / "primary_per_label_tests.csv")
    require(len(tests) == 12 and tests["p_value_bh"].between(0, 1).all(), "Per-label BH table invalid")
    calibration = _finite_csv(RESULTS / "primary_calibration_intervals.csv")
    require(calibration["n_bootstrap"].eq(2000).all(), "Calibration bootstrap count is not 2,000")
    return {"patients": 1700, "labels": 12, "oof_predictions_per_patient": 5, "models": len(ALL_PREDICTION_NAMES)}


def gcn_temporal_synthetic_gates() -> dict[str, object]:
    gcn = _finite_csv(RESULTS / "gcn_ablation" / "gcn_ablation_10_seed_full.csv")
    require(len(gcn) == 180, "GCN ablation must contain 18 configurations x 10 seeds")
    require(set(gcn["seed"]) == set(range(42, 52)), "GCN result seeds changed")
    require(gcn.groupby(["seed", "head", "loss"])["adjacency"].nunique().eq(3).all(), "Paired GCN graph variants incomplete")
    temporal = _finite_csv(RESULTS / "temporal" / "temporal_model_metrics.csv")
    require(dict(zip(temporal["horizon"], temporal["feature_count"])) == {"Admission": 91, "24 h": 94, "48 h": 97, "72 h": 100}, "Temporal feature counts changed")
    temporal_oof = np.load(RESULTS / "temporal" / "temporal_oof_predictions.npz")
    require(temporal_oof["y_true"].shape == (1700, 12), "Temporal outcome archive has the wrong shape")
    for contract in ("admission_safe_v1", "prospective_24h", "prospective_48h", "prospective_72h"):
        value = temporal_oof[contract]
        require(value.shape == (5, 1700, 12), f"Temporal OOF shape changed for {contract}: {value.shape}")
        require(np.isfinite(value).all(), f"Temporal OOF archive contains missing predictions for {contract}")
    temporal_checkpoints = list((RESULTS / "temporal").glob("temporal_*_r*_f*.npz"))
    require(len(temporal_checkpoints) == 100, "Temporal analysis must contain 4 horizons x 5 repeats x 5 folds")
    synthetic = _finite_csv(RESULTS / "synthetic" / "controlled_dependency_raw.csv")
    require(synthetic["run"].nunique() == 10, "Synthetic experiment must use ten repeats")
    require(synthetic["rho"].nunique() == 5, "Synthetic experiment must use five dependency levels")
    return {"gcn_rows": len(gcn), "temporal_horizons": 4, "synthetic_repeats": 10}


def external_gates() -> dict[str, object]:
    folder = RESULTS / "external_transportability"
    results = _finite_csv(folder / "external_transportability_results.csv")
    primary = results.query("cohort == 'Hungarian AMI registry' and endpoint == 'death_30d' and subset == 'all first events'")
    require(len(primary) == 2 and primary["n"].eq(29596).all() and primary["events"].eq(3888).all(), "30-day external counts changed")
    seven = results.query("cohort == 'Hungarian AMI registry' and endpoint == 'death_7d' and subset == 'all first events'")
    require(len(seven) == 2 and seven["events"].eq(2408).all(), "7-day external counts changed")
    complete = results.query("cohort == 'Hungarian AMI registry' and endpoint == 'death_30d' and subset == 'complete case'")
    require(len(complete) == 2 and complete["n"].eq(28477).all(), "Complete-case count changed")
    meta = json.loads((folder / "external_transportability_metadata.json").read_text(encoding="utf-8"))
    require("stress test" in json.dumps(meta).lower(), "External analysis is not labelled a stress test")
    require(meta.get("source_cv_folds") == 5, "External source CV folds changed")
    require(meta.get("logistic_max_iter") == 5000, "External logistic setting changed")
    require(meta.get("tabpfn_n_estimators") == 8, "External TabPFN ensemble size changed")
    require(meta.get("tabpfn_n_preprocessing_jobs") == 1, "External TabPFN preprocessing jobs changed")
    return {"index_events": 29596, "death_30d": 3888, "death_7d": 2408, "complete_cases": 28477}


def manuscript_gates() -> dict[str, object]:
    text = MANUSCRIPT.read_text(encoding="utf-8")
    abstract_match = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", text, re.S)
    require(abstract_match is not None, "Abstract not found")
    abstract_words = len(re.findall(r"\b[A-Za-z0-9][A-Za-z0-9'-]*\b", abstract_match.group(1)))
    require(200 <= abstract_words <= 250, f"Abstract has {abstract_words} words")
    forbidden = [
        "70/10/20", "119-feature", "119 inputs", "102--111", "0.7587", "0.7606",
        "prespecified external validation", "lacks external validation", "no external validation",
        "100% reproducible", "4/12 (fdr)", "yes (appendix)",
    ]
    lower = text.lower()
    require(not any(item in lower for item in forbidden), "Legacy claim or result remains in the clean manuscript")
    for required in (
        "admission-safe", "five repeats", "four-fold inner", "cross-cohort mortality transportability stress test",
        "29,596", "3,888", "28,477", "publicly available, de-identified data",
    ):
        require(required.lower() in lower, f"Required manuscript phrase/value missing: {required}")
    require("keepaspectratio" in text, "Image aspect-ratio protection is missing")
    body, bibliography = text.split("\\begin{thebibliography}{00}", 1)
    first_citations: list[str] = []
    for group in re.findall(r"\\cite\{([^}]+)\}", body):
        for key in map(str.strip, group.split(",")):
            if key and key not in first_citations:
                first_citations.append(key)
    bibitems = re.findall(r"\\bibitem\{([^}]+)\}", bibliography)
    require(first_citations == bibitems, "Bibliography is not in first-citation order")
    require(len(bibitems) == 29, f"Expected 29 references, found {len(bibitems)}")
    return {"abstract_words": abstract_words, "references": len(bibitems)}


def submission_artifact_gates() -> dict[str, object]:
    submission = RESULTS.parent / "submission_r1_final"
    required = {
        "Access-2026-35668_Author_Response.docx",
        "Access-2026-35668_Highlighted.pdf",
        "Access-2026-35668_Clean.pdf",
        "Access-2026-35668_Clean_LaTeX_Source.zip",
        "Access-2026-35668_Supplementary_Materials.zip",
    }
    observed = {path.name for path in submission.glob("*") if path.is_file()}
    require(required.issubset(observed), f"Submission package missing: {sorted(required - observed)}")
    supplement_zip = submission / "Access-2026-35668_Supplementary_Materials.zip"
    with zipfile.ZipFile(supplement_zip) as archive:
        names = [Path(name).name.lower() for name in archive.namelist()]
    require(not any("prediction" in name or name.endswith(".npz") for name in names), "Supplement ZIP contains patient-level predictions")
    require(not any("hungarian myocardial infarction registry" in name and name.endswith(".xlsx") for name in names), "Supplement ZIP contains the Hungarian raw workbook")
    clean = fitz.open(submission / "Access-2026-35668_Clean.pdf")
    highlighted = fitz.open(submission / "Access-2026-35668_Highlighted.pdf")
    require(len(clean) == 17 and len(highlighted) == 17, "Clean/highlighted PDF page counts differ from the audited manuscript")
    require(all(page.rect.width > 0 and page.rect.height > 0 for page in clean), "Clean PDF contains an invalid page")
    clean.close()
    highlighted.close()

    response_path = submission / "Access-2026-35668_Author_Response.docx"
    response = Document(response_path)
    response_text = "\n".join(paragraph.text for paragraph in response.paragraphs)
    require(sum(paragraph.text.startswith("Comment ") for paragraph in response.paragraphs) == 14, "Response DOCX does not contain 14 comments")
    require("Original title:" in response_text and "Revised title:" in response_text, "Response DOCX title metadata is incomplete")
    require(not any(token in response_text.lower() for token in ("placeholder", "todo", "tbd", "[insert")), "Response DOCX contains a placeholder")
    with zipfile.ZipFile(response_path) as archive:
        xml_names = set(archive.namelist())
        document_xml = archive.read("word/document.xml")
    require("word/comments.xml" not in xml_names, "Response DOCX contains comments")
    require(b"<w:ins" not in document_xml and b"<w:del" not in document_xml, "Response DOCX contains tracked changes")
    return {"submission_files": len(observed), "required_files": len(required), "pdf_pages": 17, "response_comments": 14}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--level", choices=("analysis", "submission"), default="analysis")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    report = {
        "config_and_features": config_and_feature_gates(),
        "checkpoints": checkpoint_gates(),
        "internal": internal_result_gates(),
        "gcn_temporal_synthetic": gcn_temporal_synthetic_gates(),
        "external": external_gates(),
    }
    if args.level == "submission":
        report["manuscript"] = manuscript_gates()
        report["submission_artifacts"] = submission_artifact_gates()
    path = RESULTS / f"qa_{args.level}.json"
    path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
